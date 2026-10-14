"""In-app editors for Studio files.

Edits always write back to the same file id, so the chat attachment, the Files
tab and the notebook all see the new version.

  deck      PPTX built from a StudioDocument: edit slides, add, delete, reorder
  document  PDF/DOCX/HTML/SVG built from a StudioDocument: edit title, summary, sections
  sheet     XLSX: edit tables (cells), rebuilt with formulas and native charts
  pdf       any PDF: page thumbnails, reorder and delete pages
  text      MD/TXT/JSON: plain text editing
"""

import io
from typing import Any, List, Literal, Optional

from fastapi import APIRouter, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel

from open_notebook.exports import storage
from open_notebook.exports.content import (
    Section,
    Slide,
    StudioDocument,
    StudioOptions,
    Table,
)
from open_notebook.exports.renderers import RENDERERS
from open_notebook.exports.workflow import AUDITS

router = APIRouter()

TEXT_FORMATS = {"md", "txt", "json"}
DOC_FORMATS = {"pdf", "docx", "html", "svg"}


def _item(file_id: str) -> storage.StudioFile:
    item = storage.get_file(file_id)
    if not item or not item.path.exists():
        raise HTTPException(status_code=404, detail="File not found")
    return item


def _doc(item: storage.StudioFile) -> Optional[StudioDocument]:
    return storage.load_document(item.document_id) if item.document_id else None


def _tables_from_xlsx(data: bytes) -> List[Table]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=False)
    tables = []
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(v is not None for v in r)]
        if len(rows) < 2:
            continue
        cols = [
            str(c) if c is not None else f"Column {i + 1}"
            for i, c in enumerate(rows[0])
        ]
        tables.append(
            Table(name=ws.title[:31], columns=cols, rows=[list(r) for r in rows[1:]])
        )
    return tables


def _mode(item: storage.StudioFile, doc: Optional[StudioDocument]) -> str:
    if item.format == "pptx":
        return "deck" if doc else "readonly"
    if item.format == "xlsx":
        return "sheet"
    if item.format in TEXT_FORMATS:
        return "text"
    if item.format in DOC_FORMATS:
        return "document" if doc else ("pdf" if item.format == "pdf" else "readonly")
    return "readonly"


def _audit(
    item: storage.StudioFile, data: bytes, doc: Optional[StudioDocument]
) -> List[str]:
    fn = AUDITS.get(item.format)
    if not fn or doc is None:
        return []
    try:
        return [i["message"] for i in fn(data, doc)]
    except Exception as e:
        return [f"Audit could not run: {type(e).__name__}"]


@router.get("/editor/files/{file_id}")
async def open_file(file_id: str):
    item = _item(file_id)
    doc = _doc(item)
    mode = _mode(item, doc)
    out: dict[str, Any] = {
        "file": item.model_dump(),
        "mode": mode,
        "pdf_tools": item.format == "pdf",
    }
    if doc is not None:
        out["document"] = doc.model_dump()
    if mode == "sheet":
        out["tables"] = [
            t.model_dump()
            for t in (
                doc.tables
                if doc and doc.tables
                else _tables_from_xlsx(item.path.read_bytes())
            )
        ]
    if mode == "text":
        out["text"] = item.path.read_text("utf-8", errors="replace")
    if item.format == "pdf":
        import pypdfium2 as pdfium

        pdf = pdfium.PdfDocument(item.path.read_bytes())
        out["pages"] = len(pdf)
        pdf.close()
    return out


def _save(item: storage.StudioFile, doc: Optional[StudioDocument], data: bytes):
    updated = storage.update_file(item.id, data, doc)
    issues = _audit(item, data, doc)
    storage.update_meta(item.id, audit_issues=issues)
    return {"file": updated.model_dump() if updated else None, "audit_issues": issues}


class DocumentUpdate(BaseModel):
    document: StudioDocument


@router.put("/editor/files/{file_id}/document")
async def save_document(file_id: str, body: DocumentUpdate):
    item = _item(file_id)
    if item.format not in RENDERERS:
        raise HTTPException(
            status_code=400, detail="This file type is not built from a document"
        )
    render = RENDERERS[item.format][0]
    if body.document.kind == "presentation":
        n = len(body.document.slides)
        body.document.options.slide_count = max(
            4, min(20, n + 3 if n + 3 >= 8 else n + 2)
        )
    data = render(body.document)
    return _save(item, body.document, data)


class SheetUpdate(BaseModel):
    tables: List[Table]
    title: Optional[str] = None


@router.put("/editor/files/{file_id}/sheet")
async def save_sheet(file_id: str, body: SheetUpdate):
    item = _item(file_id)
    if item.format != "xlsx":
        raise HTTPException(status_code=400, detail="Not a spreadsheet")
    doc = _doc(item) or StudioDocument(
        kind="spreadsheet", title=body.title or item.title
    )
    doc.tables = body.tables
    if body.title:
        doc.title = body.title
    data = RENDERERS["xlsx"][0](doc)
    return _save(item, doc, data)


class TextUpdate(BaseModel):
    text: str


@router.put("/editor/files/{file_id}/text")
async def save_text(file_id: str, body: TextUpdate):
    item = _item(file_id)
    if item.format not in TEXT_FORMATS:
        raise HTTPException(status_code=400, detail="Not a text file")
    updated = storage.update_file(item.id, body.text.encode("utf-8"))
    return {"file": updated.model_dump() if updated else None, "audit_issues": []}


@router.get("/editor/files/{file_id}/pages/{page}.png")
async def page_png(file_id: str, page: int, scale: float = 0.6):
    import pypdfium2 as pdfium

    item = _item(file_id)
    if item.format != "pdf":
        raise HTTPException(status_code=400, detail="Not a PDF")
    pdf = pdfium.PdfDocument(item.path.read_bytes())
    if page < 0 or page >= len(pdf):
        raise HTTPException(status_code=404, detail="Page not found")
    img = pdf[page].render(scale=max(0.2, min(scale, 2.5))).to_pil()
    buf = io.BytesIO()
    img.save(buf, "PNG")
    pdf.close()
    return Response(
        buf.getvalue(), media_type="image/png", headers={"Cache-Control": "no-store"}
    )


class PageOrder(BaseModel):
    order: List[int]


@router.post("/editor/files/{file_id}/pages")
async def reorder_pages(file_id: str, body: PageOrder):
    """Rebuild the PDF with pages in the given order; omitted pages are deleted."""
    import pypdfium2 as pdfium

    item = _item(file_id)
    if item.format != "pdf":
        raise HTTPException(status_code=400, detail="Not a PDF")
    src = pdfium.PdfDocument(item.path.read_bytes())
    if not body.order or any(i < 0 or i >= len(src) for i in body.order):
        raise HTTPException(status_code=400, detail="Invalid page order")
    out = pdfium.PdfDocument.new()
    for i in body.order:
        out.import_pages(src, [i])
    buf = io.BytesIO()
    out.save(buf)
    src.close()
    out.close()
    updated = storage.update_file(item.id, buf.getvalue())
    # The page layout no longer matches the structured document; detach it.
    storage.update_meta(item.id, document_id=None)
    return {
        "file": storage.get_file(item.id).model_dump() if updated else None,
        "audit_issues": [],
    }


class NewFile(BaseModel):
    format: Literal["pptx", "xlsx", "pdf", "docx", "md"]
    title: str = "Untitled"


@router.post("/editor/new")
async def new_file(body: NewFile):
    """Create a file from scratch; it opens in its editor and is listed in Files."""
    title = body.title.strip() or "Untitled"
    slug = (
        "".join(ch if ch.isalnum() else "-" for ch in title).strip("-")[:60]
        or "Untitled"
    )
    if body.format == "md":
        data = f"# {title}\n\n".encode("utf-8")
        item = storage.register_blank(f"{slug}.md", "md", data, "summary", title)
        return item.model_dump()
    if body.format == "pptx":
        doc = StudioDocument(
            kind="presentation",
            title=title,
            subtitle="",
            options=StudioOptions(slide_count=4),
            slides=[
                Slide(
                    layout="bullets",
                    title="Slide 1",
                    bullets=["Edit this slide"],
                ),
                Slide(layout="bullets", title="Slide 2", bullets=["Edit this slide"]),
            ],
        )
    elif body.format == "xlsx":
        doc = StudioDocument(
            kind="spreadsheet",
            title=title,
            tables=[
                Table(name="Sheet1", columns=["Item", "Value"], rows=[["Example", 0]])
            ],
        )
    else:
        doc = StudioDocument(
            kind="report",
            title=title,
            options=StudioOptions(length="short"),
            sections=[Section(heading="Introduction", body="Start writing here.")],
        )
    data = RENDERERS[body.format][0](doc)
    item = storage.register_blank(
        f"{slug}.{body.format}", body.format, data, doc.kind, title, doc
    )
    return item.model_dump()
